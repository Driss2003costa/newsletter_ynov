# -*- coding: utf-8 -*-
"""
Génère le dossier BC3 (Retour d'expérience en entreprise) d'Estelle CASTEROT
au format Word. Les zones surlignées en jaune sont à compléter par l'étudiante
(chiffres, effectif, dates, KPI) : aucun fait n'est inventé.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

JUST = WD_ALIGN_PARAGRAPH.JUSTIFY
CENTER = WD_ALIGN_PARAGRAPH.CENTER

doc = Document()

# ------------------------------------------------------------------ styles
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

for lvl, sz in [('Heading 1', 15), ('Heading 2', 13), ('Heading 3', 11.5)]:
    st = doc.styles[lvl]
    st.font.name = 'Calibri'
    st.font.size = Pt(sz)
    st.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

illustrations = []  # (label) pour la table des illustrations

# ------------------------------------------------------------------ helpers
def seg_runs(paragraph, segments):
    for seg in segments:
        if isinstance(seg, tuple):
            text, kind = seg
        else:
            text, kind = seg, 'normal'
        run = paragraph.add_run(text)
        if kind == 'ph':
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.italic = True
        elif kind == 'bold':
            run.bold = True

def para(*segments):
    p = doc.add_paragraph()
    p.alignment = JUST
    seg_runs(p, segments)
    return p

def bullet(*segments):
    p = doc.add_paragraph(style='List Bullet')
    seg_runs(p, segments)
    return p

def numbered(*segments):
    p = doc.add_paragraph(style='List Number')
    seg_runs(p, segments)
    return p

def h1(text):
    doc.add_heading(text, level=1)

def h2(text):
    doc.add_heading(text, level=2)

def placeholder(text):
    p = doc.add_paragraph()
    r = p.add_run('[À compléter : ' + text + ']')
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    r.italic = True
    return p

def figure_placeholder(caption):
    illustrations.append(caption)
    p = doc.add_paragraph()
    p.alignment = CENTER
    r = p.add_run('[ ' + caption + ' — visuel à insérer ]')
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    r.italic = True
    cap = doc.add_paragraph()
    cap.alignment = CENTER
    cr = cap.add_run(caption)
    cr.italic = True
    cr.font.size = Pt(9)

def add_table(caption, headers, rows):
    """rows : liste de listes ; chaque cellule est une chaîne ou un tuple (texte, 'ph')."""
    illustrations.append(caption)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            if isinstance(val, tuple):
                text, kind = val
                r = p.add_run(text)
                r.font.size = Pt(10)
                if kind == 'ph':
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    r.italic = True
            else:
                r = p.add_run(val)
                r.font.size = Pt(10)
    cap = doc.add_paragraph()
    cap.alignment = CENTER
    cr = cap.add_run(caption)
    cr.italic = True
    cr.font.size = Pt(9)
    doc.add_paragraph()

def page_break():
    doc.add_page_break()

# ------------------------------------------------------------------ footer page number
def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = CENTER
    run = p.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instr); run._r.append(fldChar2)

# ================================================================== PAGE DE GARDE
sec = doc.sections[0]
sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)

def gap(n=1):
    for _ in range(n):
        doc.add_paragraph()

p = doc.add_paragraph(); p.alignment = CENTER
r = p.add_run('[ Logo Ynov ]        [ Logo J\'ai vu la Vierge ]')
r.font.highlight_color = WD_COLOR_INDEX.YELLOW; r.italic = True
gap(2)

p = doc.add_paragraph(); p.alignment = CENTER
r = p.add_run('Dossier de Retour d\'Expérience en Entreprise'); r.bold = True; r.font.size = Pt(24)
p = doc.add_paragraph(); p.alignment = CENTER
r = p.add_run('Épreuve certifiante du Bloc de Compétences 3 (BC3)'); r.font.size = Pt(14)
gap(1)
p = doc.add_paragraph(); p.alignment = CENTER
r = p.add_run('Bachelor 3 Communication, Marketing et Événementiel'); r.font.size = Pt(13); r.italic = True
gap(3)

p = doc.add_paragraph(); p.alignment = CENTER
r = p.add_run('Estelle CASTEROT'); r.bold = True; r.font.size = Pt(18)
gap(2)

for label, value in [
    ('Entreprise d\'accueil', 'J\'ai vu la Vierge'),
    ('Poste occupé', 'Cheffe de Projet e-commerce'),
    ('Tutrice en entreprise', 'Camille Abela, Responsable de Projet e-commerce'),
]:
    p = doc.add_paragraph(); p.alignment = CENTER
    r = p.add_run(label + ' : '); r.bold = True; r.font.size = Pt(12)
    r2 = p.add_run(value); r2.font.size = Pt(12)

p = doc.add_paragraph(); p.alignment = CENTER
r = p.add_run('Campus Ynov '); r.bold = True; r.font.size = Pt(12)
r2 = p.add_run('['); r2.font.size = Pt(12)
r3 = p.add_run('à compléter : ville du campus'); r3.font.highlight_color = WD_COLOR_INDEX.YELLOW; r3.italic = True; r3.font.size = Pt(12)
r4 = p.add_run(']'); r4.font.size = Pt(12)

gap(3)
p = doc.add_paragraph(); p.alignment = CENTER
r = p.add_run('Année scolaire 2025 / 2026'); r.font.size = Pt(12)
page_break()

# ================================================================== REMERCIEMENTS
h1('Remerciements')
para("Avant d'entrer dans le vif de ce retour d'expérience, je souhaite adresser mes remerciements "
     "à toutes les personnes qui ont rendu cette alternance aussi riche.")
para("Je remercie tout particulièrement Camille Abela, ma tutrice et Responsable de Projet e-commerce, "
     "pour sa confiance, sa disponibilité et l'accompagnement qu'elle m'a offert tout au long de l'année. "
     "Ses briefs, ses retours et nos points hebdomadaires m'ont permis de progresser semaine après semaine "
     "et de gagner en autonomie sur un métier que je découvrais.")
para("Je remercie également la créatrice de la marque J'ai vu la Vierge ainsi que l'ensemble de l'équipe, "
     "qui m'ont accueillie avec bienveillance et m'ont fait confiance sur des sujets à forte visibilité.")
para("Mes remerciements vont enfin à l'équipe pédagogique d'Ynov et à mon référent de formation, "
     "qui m'ont donné les repères méthodologiques nécessaires pour mener à bien mes missions et "
     "pour prendre le recul attendu dans ce dossier.")
page_break()

# ================================================================== SOMMAIRE
h1('Sommaire')
para("Pour actualiser la pagination automatique, faites un clic droit sur le sommaire ci-dessous "
     "puis choisissez « Mettre à jour les champs », ou sélectionnez tout le document et appuyez sur F9.")
# champ TOC
p = doc.add_paragraph()
run = p.add_run()
fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
instr.text = 'TOC \\o "1-2" \\h \\z \\u'
fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'separate')
t = OxmlElement('w:t'); t.text = "Le sommaire s'affichera ici après mise à jour des champs (F9)."
fld3 = OxmlElement('w:fldChar'); fld3.set(qn('w:fldCharType'), 'end')
run._r.append(fld1); run._r.append(instr); run._r.append(fld2); run._r.append(t); run._r.append(fld3)
page_break()

# ================================================================== 1. INTRODUCTION
h1('1. Introduction')
para("J'ai réalisé mon alternance de troisième année de Bachelor Communication, Marketing et Événementiel "
     "chez Ynov au sein de la marque J'ai vu la Vierge. Mon rythme était de deux semaines en entreprise "
     "pour une semaine à l'école. J'y ai occupé le poste de Cheffe de Projet e-commerce, au sein du pôle "
     "digital, sous la responsabilité de ma tutrice Camille Abela, Responsable de Projet e-commerce. "
     "Cette expérience longue m'a plongée dans un univers que je connaissais peu au départ, celui du "
     "commerce en ligne, et dans lequel j'ai pris de plus en plus d'autonomie au fil des mois.")
para("J'ai vu la Vierge est une marque française qui réinvente la statuette de la Vierge Marie en objet "
     "de décoration contemporain. Elle se situe à la rencontre de l'iconographie religieuse, de l'artisanat "
     "et de la culture pop. La marque vend ses produits sur son site e-commerce propulsé par Shopify, sur "
     "une plateforme dédiée à ses revendeurs avec Odoo, et à travers un réseau de partenaires physiques "
     "reconnus comme Le Printemps, le Mucem, Maison et Objet, Boboboom ou encore l'Olympique de Marseille. "
     "Elle a aussi mené des collaborations médiatiques, notamment avec Karine Lemarchand. Ce positionnement "
     "fait à la fois sa force et sa principale difficulté de communication, car il faut réussir à proposer "
     "un objet à forte charge symbolique à un public large sans jamais trahir son identité.")
para("C'est exactement là que se situait ma mission. J'étais chargée d'animer les réseaux sociaux Instagram, "
     "TikTok et Facebook, de produire les contenus comme les posts, les stories, les shootings et les fiches "
     "produits, de participer à l'emailing et d'assurer le reporting des performances. Mon objectif était "
     "clair : développer la visibilité de la marque et soutenir ses ventes en ligne. De ce quotidien est née "
     "la question qui guide tout ce dossier.")
p = doc.add_paragraph(); p.alignment = CENTER
r = p.add_run("Comment développer la visibilité et les performances e-commerce de J'ai vu la Vierge grâce à "
              "une stratégie de communication digitale cohérente, tout en préservant une image de marque "
              "singulière et sensible ?")
r.bold = True; r.italic = True
para("Pour y répondre, je présente d'abord l'entreprise et l'analyse de son environnement, afin de poser le "
     "contexte dans lequel j'ai travaillé. Je décris ensuite mon poste et ma place dans l'équipe. J'entre "
     "alors dans le cœur du sujet en listant les constats qui ont fait naître cette problématique, puis en "
     "comparant les différentes stratégies que nous avons envisagées. Je détaille la mise en œuvre et le "
     "planning des actions menées, avant d'en mesurer les résultats à travers des indicateurs précis. Je "
     "termine en prenant du recul sur ce projet, sur ma posture professionnelle et sur les compétences de "
     "cheffe de projet que cette alternance m'a permis de développer.")
page_break()

# ================================================================== 2. ENTREPRISE
h1("2. L'entreprise et l'analyse de son environnement")

h2('2.1 Histoire et concept de la marque')
para("J'ai vu la Vierge est une marque de décoration qui revisite une figure très ancienne, la statuette de "
     "la Vierge Marie, pour en faire un objet contemporain, coloré et désirable. Là où l'objet religieux "
     "classique reste discret et traditionnel, la marque en fait une pièce de décoration assumée, à la fois "
     "spirituelle et pop, que l'on affiche chez soi comme un objet de style autant que de sens.")
placeholder("insérez ici l'année de création, le nom de la ou des personnes fondatrices, la ville et le "
            "récit d'origine de la marque, que vous retrouverez sur la page « Notre histoire » du site")
para("Ce récit d'origine est important, car il explique le ton de la marque. J'ai vu la Vierge ne se prend "
     "pas au sérieux et cultive une forme d'humour tendre, tout en respectant la charge symbolique de son "
     "sujet. C'est ce fragile équilibre entre le clin d'œil et le respect qui fait toute la singularité de la "
     "marque, et qui a guidé la plupart de mes choix de communication.")

h2('2.2 Offre, positionnement et clientèle')
para("L'offre de la marque s'articule autour de ses statuettes et des produits dérivés qui prolongent son "
     "univers. Le positionnement se situe sur un segment premium et affinitaire : on n'achète pas seulement "
     "un objet, on adhère à une histoire et à une esthétique.")
para("La clientèle est principalement féminine, avec un cœur de cible situé entre trente et soixante ans. "
     "Ce sont des femmes sensibles à la décoration, à l'objet qui a du sens et au cadeau original, qui suivent "
     "les tendances sur les réseaux sociaux et apprécient une marque qui a du caractère.")
para("Pour orienter mes contenus, je me suis appuyée sur un persona type : une femme active, urbaine, âgée "
     "d'une quarantaine d'années, attentive à la décoration de son intérieur et attirée par les objets qui "
     "racontent une histoire. Elle achète autant pour elle que pour offrir, elle est présente sur Instagram, "
     "et elle est sensible à l'authenticité d'une marque plus qu'à la simple promotion. Garder ce portrait en "
     "tête m'a aidée à choisir le bon ton, à privilégier l'émotion et le récit plutôt que le discours "
     "purement commercial, et à parler à cette communauté comme on parle à quelqu'un que l'on connaît.")
placeholder("précisez si vous le souhaitez le panier moyen, la fourchette de prix des produits et les best-sellers")

h2('2.3 Les canaux de distribution')
para("La marque a fait le choix d'une distribution à plusieurs canaux, ce qui est une vraie force pour sa "
     "visibilité comme pour ses ventes.")
bullet(("Le site e-commerce : ", 'bold'), "la boutique en ligne de la marque, propulsée par Shopify, "
        "qui constitue le cœur de l'activité digitale et le point de conversion principal.")
bullet(("La plateforme revendeurs : ", 'bold'), "un espace dédié aux professionnels et aux revendeurs, "
        "géré sous Odoo, qui permet aux partenaires de commander les produits.")
bullet(("Les partenaires physiques : ", 'bold'), "un réseau de points de vente et de lieux prestigieux "
        "comme Le Printemps, le Mucem, Maison et Objet, Boboboom ou encore l'Olympique de Marseille.")
bullet(("Les collaborations : ", 'bold'), "des opérations spéciales et médiatiques, à l'image de la "
        "collaboration menée avec Karine Lemarchand, qui offrent à la marque une exposition forte auprès "
        "de nouveaux publics.")
para("Cette présence à la fois en ligne et dans des lieux physiques emblématiques nourrit la notoriété de la "
     "marque et crédibilise son positionnement. Une grande partie de mon travail a consisté à faire vivre ces "
     "temps forts sur les canaux digitaux, pour que chaque partenariat rayonne au-delà du point de vente.")

h2("2.4 L'analyse de l'environnement externe")
para("Pour comprendre le contexte dans lequel évolue la marque, j'ai synthétisé son environnement externe à "
     "travers une grille PESTEL, qui passe en revue les grands facteurs qui influencent son activité.")
add_table(
    "Tableau 1 : Analyse PESTEL de l'environnement de J'ai vu la Vierge",
    ["Facteur", "Ce que cela implique pour la marque"],
    [
        ["Politique et réglementaire", "Contexte français de laïcité et sensibilité autour des symboles "
         "religieux dans l'espace commercial. Certains partenaires peuvent hésiter à s'associer à un objet "
         "connoté religieusement, comme je l'ai vécu avec l'Olympique de Marseille."],
        ["Économique", "Bonne dynamique du e-commerce et du marché de la décoration et du cadeau, mais forte "
         "saisonnalité liée aux fêtes de fin d'année, à la Saint Valentin et aux temps forts commerciaux."],
        ["Socioculturel", "Rapport ambivalent au religieux et goût marqué pour l'objet détourné, spirituel "
         "et pop. Les réseaux sociaux jouent un rôle central dans la diffusion des tendances déco."],
        ["Technologique", "Poids des algorithmes des plateformes sociales, essor de la vidéo courte avec les "
         "Reels et TikTok, et montée des outils d'automatisation marketing et d'intelligence artificielle "
         "appliqués au ciblage, comme la solution Kiliba."],
        ["Écologique", "Attentes croissantes des clients sur l'origine des produits et une fabrication "
         "responsable, qui deviennent des arguments de communication."],
        ["Légal", "Cadre du RGPD pour la collecte des adresses et l'emailing, droit à l'image pour les "
         "shootings et les collaborations, et obligations légales propres au e-commerce."],
    ],
)
para("Sur le plan concurrentiel, la marque évolue face à des acteurs qui exploitent le même imaginaire, "
     "notamment Sapristi et Miraculeuse. Ce qui distingue J'ai vu la Vierge de ces concurrents tient à deux "
     "atouts majeurs : une présence sur les réseaux sociaux nettement plus dynamique et travaillée, et un "
     "réseau de très bons commerçants et partenaires qui installent la marque dans des lieux prestigieux. "
     "Ces deux forces sont précisément celles sur lesquelles j'ai été amenée à travailler au quotidien.")

h2("2.5 L'analyse de l'environnement interne")
para("Sur le plan interne, J'ai vu la Vierge est une structure à taille humaine, organisée autour de sa "
     "créatrice et de plusieurs pôles qui couvrent la création, le e-commerce, la communication et la "
     "logistique. Le pôle digital, auquel j'étais rattachée, joue un rôle central puisqu'il porte à la fois "
     "l'image de la marque et une part importante de ses ventes.")
placeholder("indiquez ici l'effectif exact de l'entreprise et la répartition des équipes, par exemple "
            "création, e-commerce et digital, commercial et revendeurs, logistique")
para("La principale ressource de la marque est immatérielle : c'est la puissance de son identité et de son "
     "univers, relayée par une communauté engagée sur les réseaux sociaux. Cette force explique aussi une "
     "certaine dépendance à la portée organique de ces plateformes, un point sur lequel je reviens dans "
     "l'analyse de la problématique.")

h2('2.6 Synthèse SWOT')
para("L'ensemble de cette analyse peut se résumer dans une matrice SWOT, qui met en regard les forces et les "
     "faiblesses internes avec les opportunités et les menaces externes.")
add_table(
    "Tableau 2 : Matrice SWOT de J'ai vu la Vierge",
    ["Forces", "Faiblesses"],
    [
        ["Identité de marque forte et différenciante ; présence sociale dynamique ; réseau de partenaires "
         "prestigieux ; produit à fort pouvoir d'évocation ; bons commerçants.",
         "Dépendance à la portée organique des réseaux ; image sensible qui peut freiner certains "
         "partenariats ; notoriété encore de niche ; base de données clients à mieux exploiter."],
    ],
)
add_table(
    "Tableau 3 : SWOT, opportunités et menaces",
    ["Opportunités", "Menaces"],
    [
        ["Essor du e-commerce et de la vidéo courte ; engouement pour l'objet déco détourné ; collaborations "
         "médiatiques comme Karine Lemarchand ; fidélisation par l'emailing ; visibilité offerte par les "
         "partenaires physiques.",
         "Concurrence de Sapristi et Miraculeuse ; volatilité des algorithmes des plateformes ; sensibilité "
         "religieuse et risque d'image ; forte saisonnalité des ventes."],
    ],
)
page_break()

# ================================================================== 3. PERIMETRE
h1("3. Mon périmètre d'action")

h2('3.1 Mon poste et mes missions')
para("En tant que Cheffe de Projet e-commerce au sein du pôle digital, j'avais un rôle transversal qui "
     "touchait à la fois à la communication, à la création de contenu et au commerce en ligne. Concrètement, "
     "mes missions récurrentes étaient les suivantes.")
bullet(("L'animation des réseaux sociaux : ", 'bold'), "la publication et la programmation des posts et des "
        "stories sur Instagram, TikTok et Facebook, ainsi que l'animation de la communauté.")
bullet(("La création de contenu : ", 'bold'), "la participation aux shootings, la préparation des visuels et "
        "la déclinaison des contenus pour chaque plateforme.")
bullet(("La gestion du site : ", 'bold'), "la mise à jour du site e-commerce et la rédaction des fiches "
        "produits, pour que chaque référence soit bien présentée et donne envie d'acheter.")
bullet(("L'emailing : ", 'bold'), "la participation à la conception et à l'envoi des campagnes email à "
        "destination des clients de la marque.")
bullet(("Le reporting : ", 'bold'), "le suivi des performances et la restitution des résultats chaque "
        "semaine, afin d'ajuster les actions.")
para("Ce périmètre large est très formateur, car il oblige à passer sans cesse de la création à l'analyse, "
     "et à garder en tête l'objectif final qui est de développer la visibilité et les ventes.")

h2("3.2 Ma place dans l'organisation")
para("Au sein de l'organisation, je me situais dans le pôle digital, en lien direct et quotidien avec ma "
     "tutrice Camille Abela, qui pilotait les projets e-commerce, et sous le regard de la créatrice de la "
     "marque pour tout ce qui touchait à l'identité visuelle. L'organigramme ci-dessous situe ma position.")
figure_placeholder("Figure 1 : Organigramme simplifié et positionnement de l'alternante dans le pôle digital")
placeholder("complétez l'organigramme avec les prénoms et fonctions réels des membres de l'équipe autour de "
            "vous, par exemple la créatrice, le pôle commercial et revendeurs, le pôle logistique")

h2('3.3 Mes interactions et mes rituels de travail')
para("Mon quotidien était rythmé par des interactions régulières, à la fois en interne et vers l'extérieur.")
para("En interne, Camille était la personne qui me briefait et qui validait mes contenus avec la créatrice. "
     "Notre collaboration reposait sur des rituels clairs : un brief chaque lundi pour cadrer les priorités "
     "de la semaine, un point hebdomadaire pour faire le suivi des actions, et un reporting transmis chaque "
     "lundi pour mesurer les résultats de la semaine précédente. Cette régularité m'a beaucoup aidée à "
     "m'organiser et à prendre confiance.")
para("Vers l'extérieur, je livrais concrètement mes contenus aux clients et à la communauté de la marque, "
     "à travers les publications sur les réseaux mais aussi la relation directe par mail et par la messagerie "
     "d'Instagram et de Facebook. Je répondais aux questions, j'accompagnais les demandes et je faisais "
     "remonter les retours des clients, ce qui plaçait la communication au plus près des attentes réelles.")

h2('3.4 Deux situations vécues')
para("Deux expériences résument bien la réalité de mon périmètre, l'une plus délicate et l'autre très "
     "positive.")
para("La première a été un partenariat avec l'Olympique de Marseille. Le club hésitait à s'associer à la "
     "marque, car il craignait que l'univers religieux des statuettes soit mal interprété et paraisse trop "
     "connoté. Cette situation m'a fait comprendre très concrètement la sensibilité de l'image de la marque "
     "et l'importance de la manière dont on présente les choses, en insistant sur la dimension culturelle, "
     "artistique et populaire de l'objet plutôt que sur sa seule dimension religieuse.")
para("La seconde, à l'inverse, a été une vraie réussite : la collaboration avec Karine Lemarchand. Cette "
     "opération a offert à la marque une belle exposition et a montré la force d'une communication bien "
     "orchestrée autour d'une personnalité. Elle m'a servi d'exemple fil rouge pour comprendre comment un "
     "temps fort se prépare, se relaie sur les réseaux et se transforme en visibilité pour la marque.")
placeholder("ajoutez si possible un ou deux détails concrets sur la collaboration Karine Lemarchand, par "
            "exemple le type de produit, le format de l'opération et les retombées observées")
page_break()

# ================================================================== 4. PROBLEMATIQUE
h1('4. La problématique digitale')

h2('4.1 Les constats de départ')
para("Ma problématique n'est pas née d'une idée abstraite, mais d'une série de constats que j'ai faits en "
     "observant le fonctionnement digital de la marque. Je les ai inventoriés le plus complètement possible, "
     "car ce sont eux qui justifient toute la stratégie qui a suivi.")
bullet(("Une visibilité très dépendante de l'organique. ", 'bold'), "La notoriété de la marque reposait "
        "largement sur la portée naturelle des réseaux sociaux, donc sur des algorithmes que l'on ne "
        "maîtrise pas et qui peuvent faire varier fortement l'audience d'un contenu à l'autre.")
bullet(("Une audience engagée mais pas assez convertie. ", 'bold'), "La communauté était réactive sur les "
        "réseaux, mais ce capital d'attention ne se transformait pas suffisamment en trafic et en ventes "
        "sur le site e-commerce.")
bullet(("Une image sensible à manier avec précaution. ", 'bold'), "Le caractère religieux de l'objet, s'il "
        "fait la singularité de la marque, pouvait freiner certains partenaires et exigeait un discours très "
        "maîtrisé, comme l'a montré l'épisode avec l'Olympique de Marseille.")
bullet(("Une relation client encore peu outillée. ", 'bold'), "Le lien avec les clients passait beaucoup "
        "par la messagerie des réseaux et par le mail, sans dispositif structuré pour fidéliser et faire "
        "revenir les acheteurs.")
bullet(("Un levier emailing sous-exploité. ", 'bold'), "L'emailing, qui est un canal que l'on maîtrise "
        "totalement et qui appartient à la marque, n'était pas encore utilisé à la hauteur de son potentiel.")
placeholder("précisez ici l'état des lieux chiffré au démarrage, par exemple le nombre d'abonnés par réseau, "
            "la taille de la base d'emails, et indiquez s'il existait déjà une newsletter avant votre arrivée")

h2('4.2 Analyse de ces constats')
para("Mis bout à bout, ces constats dessinent un déséquilibre. D'un côté, la marque possède une force rare, "
     "une identité puissante et une communauté attachée. De l'autre, cette force reposait sur un canal fragile "
     "et volatil, les réseaux sociaux, et se convertissait mal en résultats commerciaux durables. Autrement "
     "dit, la marque était très visible par moments, mais cette visibilité restait dépendante des plateformes "
     "et insuffisamment transformée en trafic, en ventes et en fidélité.")
para("La sensibilité de l'image ajoutait une contrainte supplémentaire : il ne suffisait pas de communiquer "
     "plus, il fallait communiquer juste, en préservant l'équilibre entre le clin d'œil et le respect qui fait "
     "l'âme de la marque. Toute solution devait donc répondre à un double impératif de performance et de "
     "cohérence d'image.")
para("Cet enjeu est d'autant plus important que les deux forces de la marque, sa communauté et son réseau de "
     "partenaires, dépendent l'une comme l'autre de la maîtrise de cette image. Un contenu mal calibré peut "
     "faire fuir une partie de l'audience, et un partenaire hésitant peut renoncer si l'univers religieux "
     "prend trop de place, comme l'a montré la réticence de l'Olympique de Marseille. La communication n'est "
     "donc pas seulement un levier de croissance pour la marque, c'est aussi un poste sensible qui protège ou "
     "fragilise sa réputation. C'est pour cette raison que j'ai abordé chaque action non pas isolément, mais "
     "comme une pièce d'une stratégie d'ensemble qui devait rester cohérente du premier post jusqu'au dernier "
     "email.")

h2('4.3 Formulation de la problématique')
para("De cette analyse découle la problématique qui a structuré mon travail tout au long de l'année.")
p = doc.add_paragraph(); p.alignment = CENTER
r = p.add_run("Comment développer la visibilité et les performances e-commerce de J'ai vu la Vierge grâce à "
              "une stratégie de communication digitale cohérente, tout en préservant une image de marque "
              "singulière et sensible ?")
r.bold = True; r.italic = True
page_break()

# ================================================================== 5. SCENARIOS
h1('5. Les scénarios de stratégie envisagés')

h2("5.1 Les objectifs de l'entreprise")
para("Avant de comparer les stratégies possibles, il faut rappeler les objectifs que la marque m'avait "
     "fixés, car ce sont eux qui servent de boussole pour juger de la pertinence de chaque scénario.")
bullet("Développer la visibilité et la notoriété de la marque.")
bullet("Augmenter le trafic et les ventes sur le site e-commerce.")
bullet("Élargir l'audience sans diluer ni trahir l'identité de la marque.")
bullet("Renforcer la relation et la fidélité des clients existants.")

h2('5.2 Les scénarios comparés')
para("À partir de ces objectifs, j'ai raisonné autour de trois scénarios possibles pour la stratégie de "
     "communication digitale.")
para(("Scénario A, tout miser sur l'organique social. ", 'bold'),
     "Ce scénario consiste à concentrer tous les efforts sur l'animation des réseaux sociaux, avec beaucoup "
     "de contenu et un travail de communauté approfondi. Il est peu coûteux et très cohérent avec l'ADN de "
     "la marque, mais il maintient la dépendance aux algorithmes et convertit mal l'audience en ventes.")
para(("Scénario B, industrialiser l'emailing et le CRM. ", 'bold'),
     "Ce scénario mise sur l'automatisation marketing, par exemple avec une solution comme Kiliba, pour "
     "convertir et fidéliser grâce à des emails ciblés déclenchés selon le comportement des clients. Il est "
     "efficace sur la conversion, mais il demande un budget d'abonnement et un contrôle attentif du design "
     "pour rester fidèle à l'identité de la marque.")
para(("Scénario C, une stratégie mixte et cohérente. ", 'bold'),
     "Ce scénario combine l'animation des réseaux pour la visibilité, l'emailing pour la conversion et la "
     "fidélisation, et les temps forts comme les collaborations pour créer de l'événement. C'est le scénario "
     "le plus complet, celui qui répond à la fois aux objectifs de visibilité et de ventes tout en gardant la "
     "main sur l'image.")

h2('5.3 Tableau comparatif')
add_table(
    "Tableau 4 : Comparaison des trois scénarios de stratégie",
    ["Critère", "Scénario A\nTout organique", "Scénario B\nEmailing / CRM", "Scénario C\nStratégie mixte"],
    [
        ["Coût", "Faible", "Moyen (abonnement)", "Maîtrisé et progressif"],
        ["Délai de mise en place", "Court", "Moyen", "Progressif"],
        ["Cohérence avec l'image", "Très forte", "À surveiller", "Forte et maîtrisée"],
        ["Flexibilité", "Élevée", "Moyenne", "Élevée"],
        ["Impact sur la visibilité", "Fort", "Faible", "Fort"],
        ["Impact sur les ventes", "Limité", "Fort", "Fort"],
    ],
)

h2('5.4 Le scénario retenu et sa justification')
para("Le scénario retenu est le scénario C, la stratégie mixte. C'est le seul qui répond à l'ensemble des "
     "objectifs de la marque en même temps : il entretient la visibilité par les réseaux, il transforme cette "
     "visibilité en ventes grâce à l'emailing, et il garde la maîtrise de l'image en s'appuyant sur des temps "
     "forts soigneusement préparés. Il évite le principal défaut du scénario A, qui plafonne sur la conversion, "
     "et celui du scénario B, qui peut fragiliser l'identité visuelle s'il repose uniquement sur des modèles "
     "standardisés.")
para("À l'intérieur de ce scénario, le levier emailing posait une question d'outil que j'ai instruite à part, "
     "car plusieurs solutions étaient possibles.")
add_table(
    "Tableau 5 : Comparaison des solutions d'emailing envisagées",
    ["Solution", "Atouts", "Limites"],
    [
        ["Kiliba", "Ciblage piloté par l'intelligence artificielle selon le comportement des utilisateurs, "
         "ce qui permet de mieux cibler les clients ; intégration directe à Shopify ; mise en place rapide.",
         "Personnalisation graphique plus limitée ; coût d'abonnement ; moins de contrôle sur l'identité "
         "visuelle des emails."],
        ["Shopify natif (Shopify Email)", "Intégré à la boutique et simple d'utilisation ; peu coûteux.",
         "Fonctionnalités et design limités ; canal peu différenciant qui ne valorise pas assez l'univers "
         "de la marque."],
        ["Outil sur mesure (studio interne)", "Contrôle total de l'identité visuelle grâce à une direction "
         "artistique dédiée ; adaptable aux temps forts comme un vernissage ou un programme ; gratuit à "
         "l'usage.",
         "Nécessite du développement et de la maintenance."],
    ],
)
para("Le raisonnement a conduit à privilégier Kiliba pour la partie automatisée et ciblée, parce que son "
     "intelligence artificielle centrée sur le comportement des utilisateurs permettait de mieux cibler les "
     "envois, et parce que son intégration directe à Shopify simplifiait beaucoup la mise en place. En "
     "parallèle, un studio d'emailing sur mesure a été développé pour les communications les plus premium, "
     "celles où l'identité visuelle de la marque doit être parfaitement respectée. Je détaille cet outil dans "
     "la mise en œuvre.")
para("Ce choix est cohérent avec les objectifs de l'entreprise, et c'est le critère qui a primé. La stratégie "
     "mixte sert directement la visibilité grâce aux réseaux, les ventes grâce à l'emailing et la conversion, "
     "et la fidélité grâce à une relation client mieux outillée, sans jamais sacrifier l'identité de la marque. "
     "J'ai aussi tenu compte des risques de chaque option. Le principal risque du scénario retenu est la "
     "charge de travail, puisqu'il faut animer plusieurs canaux en même temps. Je l'ai atténué en planifiant "
     "rigoureusement les tâches sur Trello et en concentrant les efforts sur les temps forts à plus fort "
     "impact. Le second risque, celui de diluer l'image en multipliant les prises de parole, a été maîtrisé "
     "par le circuit de validation systématique avec la créatrice et ma tutrice. Ce raisonnement par les "
     "risques m'a permis de défendre le scénario retenu de manière argumentée et non par simple intuition.")
page_break()

# ================================================================== 6. MISE EN OEUVRE
h1('6. Mise en œuvre et planning')

h2('6.1 Ma méthodologie de gestion de projet')
para("Pour piloter cette stratégie, j'ai travaillé de façon structurée, en m'appuyant sur des méthodes de "
     "gestion de projet simples mais efficaces. Mon outil central était Trello, sur lequel je planifiais "
     "l'ensemble des tâches à réaliser sous forme de tableau, en suivant l'avancement de chaque contenu, de "
     "l'idée jusqu'à la publication. Cette organisation en colonnes, proche d'une logique kanban, me permettait "
     "de visualiser d'un coup d'œil ce qui était à faire, en cours et terminé.")
para("Chaque création suivait un circuit de validation clair. Je préparais les visuels et les contenus, puis "
     "ils étaient validés par la créatrice de la marque et par ma tutrice avant diffusion. Ce principe de "
     "validation en deux temps garantissait la cohérence avec l'identité de la marque et me protégeait des "
     "faux pas, surtout sur un univers aussi sensible. Nos rituels hebdomadaires, le brief du lundi, le point "
     "de suivi et le reporting, complétaient cette méthode en donnant un rythme régulier au projet.")

h2('6.2 Le planning des actions')
para("Les actions se sont déployées progressivement sur l'année, selon un rétroplanning que je résume "
     "ci-dessous. Les périodes exactes sont à ajuster selon votre calendrier réel.")
add_table(
    "Tableau 6 : Rétroplanning simplifié des actions sur l'année",
    ["Période", "Phase", "Actions principales"],
    [
        [("[à préciser]", 'ph'), "Prise en main et observation", "Découverte de la marque, de ses outils "
         "(Shopify, Odoo, Trello) et de sa ligne éditoriale ; premiers contenus accompagnés."],
        [("[à préciser]", 'ph'), "Montée en autonomie", "Animation régulière des réseaux, production des "
         "contenus et des fiches produits, mise en place des rituels de reporting."],
        [("[à préciser]", 'ph'), "Temps forts et collaborations", "Préparation et relais des opérations "
         "spéciales, dont la collaboration avec Karine Lemarchand ; travail sur les partenariats."],
        [("[à préciser]", 'ph'), "Structuration de l'emailing", "Choix et mise en place des solutions "
         "d'emailing, création des premiers modèles et campagnes."],
        [("[à préciser]", 'ph'), "Optimisation et bilan", "Analyse des performances, ajustements de la ligne "
         "éditoriale et bilan de la stratégie."],
    ],
)
figure_placeholder("Figure 2 : Diagramme de Gantt du projet sur l'année d'alternance")

h2('6.3 Le levier des réseaux sociaux')
para("Le premier levier de la stratégie a été l'animation des réseaux sociaux, qui reste le point de contact "
     "principal entre la marque et son public. J'ai travaillé une ligne éditoriale fidèle au ton de la marque, "
     "en alternant les contenus produits, les coulisses, les temps forts et les formats plus spontanés. J'ai "
     "particulièrement investi la vidéo courte, avec les Reels et TikTok, car c'est le format qui offre "
     "aujourd'hui la meilleure portée organique et qui permet de toucher de nouveaux publics.")
para("Pour structurer cette animation, j'ai raisonné en grandes familles de contenus qui revenaient "
     "régulièrement. Les contenus produits mettaient en valeur les statuettes et donnaient envie de les "
     "acquérir. Les contenus coulisses montraient l'envers du décor, les shootings et la fabrication, pour "
     "créer de la proximité. Les contenus événementiels relayaient les temps forts et les partenariats. Enfin, "
     "les contenus plus spontanés et inspirés des tendances servaient à capter de nouveaux publics et à "
     "entretenir la présence de la marque. Cet équilibre entre plusieurs types de contenus permettait de "
     "rester intéressante sans lasser, et de nourrir à la fois la notoriété et la préférence de marque.")
para("La collaboration avec Karine Lemarchand illustre bien cette logique : autour de ce temps fort, j'ai "
     "décliné des contenus adaptés à chaque plateforme pour maximiser l'exposition et prolonger l'événement "
     "au delà du simple partenariat. C'est ce travail de mise en récit qui transforme une opération en "
     "visibilité durable pour la marque.")

h2("6.4 Le levier de l'emailing et l'outil sur mesure")
para("Le second levier, l'emailing, visait à reprendre la main sur la relation client et à convertir "
     "l'audience en ventes. Comme expliqué plus haut, la solution Kiliba a été privilégiée pour sa capacité "
     "à cibler les envois grâce à l'intelligence artificielle et pour son intégration à Shopify.")
para("En complément, un studio d'emailing sur mesure a été conçu pour les communications où l'identité "
     "visuelle prime, comme les invitations et les temps forts. Cet outil repose sur une direction artistique "
     "de type musée, sobre et élégante, avec des modèles prêts à l'emploi : une invitation de type vernissage, "
     "un email de bienvenue et un programme du mois. L'expéditeur choisit un modèle, renseigne les contenus, "
     "prévisualise le rendu comme une œuvre encadrée, puis envoie la campagne. Ce studio garantit que chaque "
     "email respecte parfaitement l'univers de la marque, ce qu'un outil standardisé ne permet pas toujours. "
     "Des captures d'écran de l'outil figurent en annexe.")

h2('6.5 Les outils utilisés')
para("Au fil du projet, je me suis appuyée sur un écosystème d'outils complémentaires.")
bullet(("Shopify : ", 'bold'), "gestion du site e-commerce, des fiches produits et des données de vente.")
bullet(("Odoo : ", 'bold'), "gestion de la relation avec les revendeurs.")
bullet(("Trello : ", 'bold'), "planification des tâches et suivi de l'avancement des contenus.")
bullet(("Les plateformes sociales et leurs outils de gestion : ", 'bold'), "Instagram, TikTok et Facebook, "
        "avec la programmation et le suivi des publications.")
bullet(("Kiliba : ", 'bold'), "emailing ciblé et automatisé, intégré à Shopify.")
bullet(("Le studio d'emailing sur mesure : ", 'bold'), "pour les campagnes premium et les temps forts.")
placeholder("ajoutez les outils de création que vous utilisiez pour les visuels, par exemple Canva ou la "
            "suite Adobe, ainsi que tout autre outil de reporting")
page_break()

# ================================================================== 7. KPI
h1('7. Les indicateurs de performance')

h2('7.1 Les indicateurs définis')
para("Pour piloter la stratégie et mesurer son efficacité, j'ai suivi des indicateurs de performance, ou KPI, "
     "répartis selon les grandes étapes du parcours client, de la notoriété jusqu'à la fidélisation. J'ai "
     "choisi de ne pas regarder un seul chiffre, mais de suivre l'ensemble de l'entonnoir de conversion. En "
     "effet, un grand nombre d'abonnés ne sert à rien s'il ne génère pas de trafic, et du trafic ne vaut que "
     "s'il se transforme en ventes puis en clients fidèles. Relier ces indicateurs entre eux permet de "
     "comprendre où se situent les points forts et les points de fuite du parcours, et donc d'agir au bon "
     "endroit. Le tableau ci-dessous présente les indicateurs retenus et l'outil qui permet de les mesurer.")
add_table(
    "Tableau 7 : Les indicateurs de performance suivis",
    ["Objectif", "Indicateurs", "Outil de mesure"],
    [
        ["Notoriété et visibilité", "Nombre d'abonnés par réseau, portée et impressions",
         "Meta Business Suite, TikTok Analytics"],
        ["Engagement", "Taux d'engagement, enregistrements, partages, commentaires",
         "Outils natifs des plateformes"],
        ["Trafic", "Sessions du site, part du trafic issu des réseaux et de l'emailing",
         "Shopify Analytics"],
        ["Conversion", "Taux de conversion, chiffre d'affaires en ligne, panier moyen",
         "Shopify"],
        ["Emailing", "Taux d'ouverture, taux de clic, taux de désabonnement",
         "Kiliba, studio d'emailing"],
        ["Fidélisation", "Taux de réachat, valeur vie client",
         "Shopify"],
    ],
)

h2('7.2 Les résultats obtenus')
para("J'ai suivi ces indicateurs dans mes reportings hebdomadaires afin de mesurer l'évolution avant et après "
     "les actions menées. Le tableau ci-dessous est prêt à recevoir vos chiffres réels, ce qui rendra cette "
     "partie particulièrement convaincante pour le jury.")
add_table(
    "Tableau 8 : Résultats avant et après la mise en œuvre de la stratégie",
    ["Indicateur", "Avant", "Après", "Objectif"],
    [
        ["Abonnés Instagram", ("[à compléter]", 'ph'), ("[à compléter]", 'ph'), ("[à compléter]", 'ph')],
        ["Abonnés TikTok", ("[à compléter]", 'ph'), ("[à compléter]", 'ph'), ("[à compléter]", 'ph')],
        ["Abonnés Facebook", ("[à compléter]", 'ph'), ("[à compléter]", 'ph'), ("[à compléter]", 'ph')],
        ["Taux d'engagement moyen", ("[à compléter]", 'ph'), ("[à compléter]", 'ph'), ("[à compléter]", 'ph')],
        ["Sessions du site", ("[à compléter]", 'ph'), ("[à compléter]", 'ph'), ("[à compléter]", 'ph')],
        ["Taux de conversion", ("[à compléter]", 'ph'), ("[à compléter]", 'ph'), ("[à compléter]", 'ph')],
        ["Taux d'ouverture email", ("[à compléter]", 'ph'), ("[à compléter]", 'ph'), ("[à compléter]", 'ph')],
        ["Taux de clic email", ("[à compléter]", 'ph'), ("[à compléter]", 'ph'), ("[à compléter]", 'ph')],
    ],
)
figure_placeholder("Figure 3 : Évolution de la visibilité et des performances avant et après les actions")

h2('7.3 Lecture des résultats')
para("Au delà des chiffres bruts, ce qui compte est la tendance qu'ils dessinent. L'enjeu de cette lecture "
     "est de relier chaque évolution à une action précise, pour montrer que les résultats ne sont pas dus au "
     "hasard mais bien au travail mené sur les deux leviers de la stratégie.")
placeholder("commentez ici vos chiffres réels : quelles actions ont le mieux fonctionné, quels temps forts "
            "ont généré le plus de visibilité et de ventes, et où se situent encore les marges de progrès")
page_break()

# ================================================================== 8. PRISE DE RECUL
h1('8. Prise de recul')

h2('8.1 Bilan critique du projet')
para("Avec du recul, je porte un regard lucide sur cette année. Les réussites sont réelles. J'ai gagné en "
     "autonomie de façon très nette, j'ai appris à faire vivre une marque sur les réseaux et à préparer des "
     "temps forts, et j'ai contribué à des opérations dont je suis fière, en particulier la collaboration "
     "avec Karine Lemarchand.")
para("Il y a aussi eu des difficultés, et c'est normal. La plus marquante a été la sensibilité de l'image de "
     "la marque, illustrée par l'hésitation de l'Olympique de Marseille à s'associer à un univers religieux. "
     "Cette situation m'a appris qu'en communication, la manière de présenter les choses compte autant que le "
     "message lui-même.")
placeholder("ajoutez ici deux ou trois difficultés ou erreurs concrètes que vous avez rencontrées, et surtout "
            "ce que vous feriez différemment aujourd'hui : c'est exactement ce que le jury attend dans une "
            "prise de recul")
para("Une autre difficulté, plus discrète, a été d'apprendre à jongler entre des tâches très différentes dans "
     "une même journée, en passant de la création d'un visuel à l'analyse de chiffres, puis à une réponse "
     "client. Au début, cette diversité me dispersait un peu. J'ai appris à la structurer grâce à la "
     "planification et aux rituels de la semaine, et c'est justement cette capacité à tenir plusieurs sujets "
     "de front qui fait, selon moi, une bonne cheffe de projet.")
para("Si je devais recommencer, je structurerais sans doute plus tôt le levier de l'emailing et la collecte "
     "des données clients, car c'est le canal qui appartient vraiment à la marque et qui offre le plus de "
     "marge de progression pour transformer la visibilité en ventes durables. Je mettrais aussi en place dès "
     "le départ un tableau de suivi des indicateurs plus complet, pour mesurer l'impact de chaque action de "
     "façon encore plus fine et pouvoir défendre mes résultats avec des preuves chiffrées.")

h2("8.2 L'évolution de ma posture professionnelle")
para("Cette alternance m'a fait évoluer sur le plan professionnel, bien au delà des compétences techniques. "
     "La transformation la plus importante est l'autonomie. Au début, j'avais besoin d'être guidée sur chaque "
     "contenu ; à la fin, je proposais, je planifiais et j'assumais mes choix, tout en respectant les "
     "validations nécessaires.")
para("J'ai aussi beaucoup progressé en communication, en apprenant à échanger avec ma tutrice, avec la "
     "créatrice et avec des partenaires extérieurs, chacun ayant ses attentes. Enfin, j'ai développé une "
     "vraie rigueur et un sens de l'organisation, indispensables pour gérer plusieurs sujets en parallèle et "
     "tenir un rythme de reporting régulier.")
para("Trois exemples concrets résument bien cette évolution. Au début, lorsqu'il fallait répondre à des "
     "commentaires négatifs sur les réseaux, je demandais systématiquement de l'aide, car l'exercice n'était "
     "pas simple et je craignais de mal formuler ma réponse. Avec le temps, j'ai appris à gérer ces messages "
     "seule, en gardant le bon ton et en protégeant l'image de la marque. Le reporting a été un autre cap. "
     "Prendre la parole pour présenter mes résultats m'intimidait et je l'appréhendais chaque semaine, mais à "
     "force de le pratiquer lors de nos points hebdomadaires, j'ai gagné en aisance, au point d'être "
     "aujourd'hui à l'aise pour exposer et défendre mon travail. Enfin, j'ai appris à décider seule quand la "
     "situation l'exigeait et à être force de proposition, notamment en suggérant des idées de posts et de "
     "contenus plutôt que d'attendre qu'on me les dicte. Ce sont ces petits pas, répétés semaine après "
     "semaine, qui m'ont fait passer d'une posture d'exécutante à une posture de cheffe de projet.")

h2('8.3 Les compétences de cheffe de projet développées')
para("Au fil de l'année, j'ai développé un ensemble de compétences qui correspondent à celles d'une cheffe "
     "de projet digital, à la fois techniques et humaines.")
add_table(
    "Tableau 9 : Les compétences développées durant l'alternance",
    ["Compétence", "Comment je l'ai développée"],
    [
        ["Gestion de projet et planification", "Pilotage des tâches sur Trello, respect des délais et "
         "coordination des validations."],
        ["Community management et création de contenu", "Animation quotidienne d'Instagram, TikTok et "
         "Facebook, production de posts, stories et vidéos courtes."],
        ["Emailing et relation client", "Participation aux campagnes email, réflexion sur les outils Kiliba "
         "et le studio sur mesure."],
        ["Analyse de performance et reporting", "Suivi hebdomadaire des KPI et restitution des résultats."],
        ["Maîtrise des outils e-commerce", "Utilisation de Shopify pour le site et les fiches produits, et "
         "d'Odoo pour les revendeurs."],
        ["Compétences humaines", "Autonomie, communication, adaptabilité, rigueur et gestion des priorités."],
    ],
)

para("Ce qui me semble le plus important, c'est que ces compétences ne sont pas restées théoriques. Chacune "
     "a été mise à l'épreuve sur des situations réelles, avec des enjeux visibles pour la marque. La gestion "
     "de projet s'est jouée sur des délais concrets, la création de contenu sur des publications vues par des "
     "milliers de personnes, et la relation client sur des échanges directs par messagerie. C'est cette "
     "confrontation au réel qui a transformé des notions apprises à l'école en réflexes professionnels, et qui "
     "me donne aujourd'hui la légitimité de me projeter vers un rôle de cheffe de projet.")
h2('8.4 Ma projection professionnelle')
para("Cette expérience a confirmé mon envie de continuer dans cette voie. À l'issue de mon Bachelor, je "
     "souhaite poursuivre mes études en mastère afin d'approfondir mes compétences et, à terme, piloter des "
     "projets de plus grande ampleur. Cette alternance m'a donné le goût du pilotage de projet et la confiance "
     "nécessaire pour viser des responsabilités de cheffe de projet à part entière.")
page_break()

# ================================================================== 9. CONCLUSION
h1('9. Conclusion')
para("Au terme de cette alternance chez J'ai vu la Vierge, je peux répondre à la question qui a guidé ce "
     "dossier. Développer la visibilité et les performances e-commerce d'une marque à l'identité aussi forte "
     "que sensible suppose une stratégie de communication digitale à la fois ambitieuse et respectueuse de "
     "son univers. C'est en combinant l'animation des réseaux sociaux, l'emailing et les temps forts, tout en "
     "gardant une maîtrise constante de l'image, que l'on transforme une communauté engagée en résultats "
     "commerciaux durables.")
para("Sur le plan personnel, cette année m'a fait grandir. J'ai découvert le e-commerce, un terrain que je ne "
     "connaissais pas, j'ai gagné en autonomie et j'ai développé des compétences concrètes de cheffe de projet. "
     "Les difficultés rencontrées, comme la sensibilité de l'image de la marque, ont été autant d'occasions "
     "d'apprendre à communiquer avec justesse.")
para("Cette expérience conforte mon projet professionnel. Elle me donne envie de poursuivre en mastère et de "
     "continuer à piloter des projets digitaux, avec la conviction qu'une bonne stratégie n'est jamais "
     "seulement une affaire de chiffres, mais aussi de sens et de cohérence.")
page_break()

# ================================================================== GLOSSAIRE
h1('Glossaire')
glossaire = [
    ("A/B test", "méthode qui consiste à comparer deux versions d'un contenu pour retenir la plus performante."),
    ("Benchmark", "analyse comparative des pratiques des concurrents ou du marché."),
    ("Brief", "document ou échange qui cadre les objectifs et les attentes d'une action ou d'un contenu."),
    ("CRM", "gestion de la relation client, c'est-à-dire l'ensemble des outils et méthodes pour suivre et "
     "fidéliser les clients."),
    ("E-commerce", "vente de produits ou de services sur internet."),
    ("KPI", "indicateur clé de performance, chiffre qui permet de mesurer l'atteinte d'un objectif."),
    ("Kanban", "méthode d'organisation visuelle des tâches par colonnes, du à faire au terminé."),
    ("Persona", "portrait type d'un client cible, utilisé pour orienter la communication."),
    ("Portée (reach)", "nombre de personnes uniques ayant vu un contenu."),
    ("Reporting", "restitution régulière des résultats et des indicateurs."),
    ("Rétroplanning", "planning construit à rebours à partir de la date finale d'un projet."),
    ("Taux d'engagement", "rapport entre les interactions et l'audience d'un contenu."),
    ("Taux de conversion", "part des visiteurs qui réalisent l'action attendue, par exemple un achat."),
    ("Taux d'ouverture", "part des destinataires qui ouvrent un email."),
    ("Template", "modèle prêt à l'emploi, ici pour concevoir rapidement un email cohérent."),
]
for term, desc in glossaire:
    p = doc.add_paragraph()
    p.alignment = JUST
    r = p.add_run(term + " : "); r.bold = True
    p.add_run(desc)
page_break()

# ================================================================== TABLE DES ILLUSTRATIONS
h1('Table des illustrations')
para("La liste ci-dessous récapitule les tableaux et figures du dossier, dans leur ordre d'apparition.")
for item in illustrations:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item)
page_break()

# ================================================================== ANNEXES
h1('Annexes')
para("Les annexes rassemblent les éléments concrets qui appuient ce dossier. Insérez ci-dessous vos captures "
     "et documents réels.")
annexes = [
    "Annexe 1 : captures d'écran du studio d'emailing sur mesure (tableau de bord, composition, aperçu d'un "
    "modèle vernissage).",
    "Annexe 2 : exemples de publications et de visuels réalisés pour les réseaux sociaux.",
    "Annexe 3 : capture du tableau Trello de suivi des tâches.",
    "Annexe 4 : exemple de reporting hebdomadaire.",
    "Annexe 5 : extraits de statistiques (réseaux sociaux, site, emailing).",
    "Annexe 6 : visuels de la collaboration avec Karine Lemarchand.",
]
for a in annexes:
    p = doc.add_paragraph()
    r = p.add_run(a); r.bold = True
    placeholder_p = doc.add_paragraph()
    rr = placeholder_p.add_run('[ document ou capture à insérer ]')
    rr.font.highlight_color = WD_COLOR_INDEX.YELLOW; rr.italic = True

# ------------------------------------------------------------------ footer / meta
add_page_number(doc.sections[0])
doc.core_properties.title = "Dossier BC3 - Retour d'expérience - Estelle CASTEROT"
doc.core_properties.author = "Estelle CASTEROT"

out = "/home/user/newsletter_ynov/dossier-bc3/Dossier_BC3_Estelle_CASTEROT.docx"
doc.save(out)
print("OK ->", out)
